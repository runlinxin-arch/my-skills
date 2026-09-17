#!/usr/bin/env python3
"""Generate test fixtures for the file-to-markdown skill evals."""
from pathlib import Path

FIXTURES = Path(__file__).parent / "fixtures"
FIXTURES.mkdir(parents=True, exist_ok=True)

# --- docx ---
from docx import Document

doc = Document()
doc.add_heading("2026 年第二季度总结", level=1)
doc.add_heading("一、业绩概览", level=2)
doc.add_paragraph("本季度营收达到 1.2 亿元，同比增长 18%。主要产品线的表现如下：")
table = doc.add_table(rows=4, cols=3)
table.style = "Table Grid"
for i, row in enumerate([
    ["产品线", "营收（万元）", "同比增长"],
    ["旗舰产品", "6800", "22%"],
    ["标准产品", "3900", "11%"],
    ["定制服务", "1300", "25%"],
]):
    for j, text in enumerate(row):
        table.rows[i].cells[j].text = text
doc.add_heading("二、下季度计划", level=2)
doc.add_paragraph("重点推进海外市场的本地化部署，目标新增 3 个区域节点。")
doc.save(FIXTURES / "季度总结.docx")

# --- pptx ---
from pptx import Presentation

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "星辰计划产品方案"
slide.placeholders[1].text = "2026 年 8 月 · 产品部"
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "核心功能"
slide.placeholders[1].text = "智能文档解析\n多格式一键转换\n批量处理与云端同步"
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "里程碑"
slide.placeholders[1].text = "Q3 完成内测\nQ4 正式上线\n2027 Q1 海外版发布"
prs.save(FIXTURES / "产品方案.pptx")

# --- xlsx ---
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "月度销售"
for row in [
    ["月份", "销售额（万元）", "订单数"],
    ["一月", 320, 1450],
    ["二月", 285, 1320],
    ["三月", 410, 1890],
]:
    ws.append(row)
wb.save(FIXTURES / "销售数据.xlsx")

# --- pdf ---
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

pdfmetrics.registerFont(TTFont("PingFang", "/System/Library/Fonts/Supplemental/Songti.ttc", subfontIndex=0))
c = canvas.Canvas(str(FIXTURES / "团队白皮书.pdf"), pagesize=A4)
c.setFont("PingFang", 18)
c.drawString(72, 780, "团队工程白皮书")
c.setFont("PingFang", 12)
c.drawString(72, 740, "第一章 协作原则")
c.drawString(72, 720, "我们坚持小步快跑、持续交付的工程文化，所有变更都要经过代码评审。")
c.drawString(72, 690, "第二章 技术选型")
c.drawString(72, 670, "后端以 Python 和 Go 为主，数据管线采用 Airflow 调度，存储使用 PostgreSQL。")
c.save()

# --- epub ---
from ebooklib import epub

book = epub.EpubBook()
book.set_title("小书")
book.set_language("zh")
ch1 = epub.EpubHtml(title="第一章 起点", file_name="ch1.xhtml", lang="zh")
ch1.content = "<h1>第一章 起点</h1><p>清晨的雾气还没有散尽，林舟已经站在了渡口。他此行的目的，是寻找传说中的古籍。</p>"
ch2 = epub.EpubHtml(title="第二章 山中", file_name="ch2.xhtml", lang="zh")
ch2.content = "<h1>第二章 山中</h1><p>山路崎岖，走了整整三天。第四天黄昏，他终于看见了那座藏在松林里的书院。</p>"
book.add_item(ch1)
book.add_item(ch2)
book.toc = [ch1, ch2]
book.spine = ["nav", ch1, ch2]
book.add_item(epub.EpubNcx())
book.add_item(epub.EpubNav())
epub.write_epub(str(FIXTURES / "小书.epub"), book)

for f in sorted(FIXTURES.iterdir()):
    print(f.name, f.stat().st_size, "bytes")
